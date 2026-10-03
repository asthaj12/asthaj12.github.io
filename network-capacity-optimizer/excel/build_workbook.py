"""Builds excel/constrained_allocation_solver.xlsx — Model 3 laid out for Excel Solver,
with the Solver dialog PRE-LOADED (objective, changing cells, constraints, Simplex LP)
so Data > Solver > Solve works immediately. Also writes the PuLP answer alongside for checking."""
import os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.workbook.defined_name import DefinedName

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'models'))
import constrained_allocation as ca

wb = Workbook()
ws = wb.active
ws.title = 'Allocation'
blue = PatternFill('solid', fgColor='DCEBFA'); yellow = PatternFill('solid', fgColor='FFF4C2'); grey = PatternFill('solid', fgColor='F2F2F2')
bold = Font(bold=True); thin = Side(style='thin', color='BBBBBB'); box = Border(top=thin, bottom=thin, left=thin, right=thin)

ws['A1'] = 'Constrained Supply Allocation — Excel Solver model'; ws['A1'].font = Font(bold=True, size=14)
ws['A2'] = 'Supply is 30% below demand. Yellow cells are the decision. Run: Data > Solver > Solve (settings are pre-loaded).'
ws['A4'] = 'Supply available'; ws['B4'] = ca.SUPPLY; ws['B4'].fill = blue; ws['A4'].font = bold

hdr = ['Channel', 'Demand', 'Margin $/unit', 'Fill floor', 'Min units', 'ALLOCATION', 'Fill %', 'Margin $', 'PuLP answer']
for j, h in enumerate(hdr, 1):
    c = ws.cell(row=6, column=j, value=h); c.font = bold; c.fill = grey; c.border = box; c.alignment = Alignment(horizontal='center')
pulp_ans = ca.solve()['allocation']
r0 = 7
for i, (name, (d, mgn, f)) in enumerate(ca.CHANNELS.items()):
    r = r0 + i
    ws.cell(row=r, column=1, value=name)
    ws.cell(row=r, column=2, value=d).fill = blue
    ws.cell(row=r, column=3, value=mgn).fill = blue
    ws.cell(row=r, column=4, value=f).fill = blue; ws.cell(row=r, column=4).number_format = '0%'
    ws.cell(row=r, column=5, value=f'=B{r}*D{r}')
    ws.cell(row=r, column=6, value=0).fill = yellow
    ws.cell(row=r, column=7, value=f'=IFERROR(F{r}/B{r},0)').number_format = '0%'
    ws.cell(row=r, column=8, value=f'=F{r}*C{r}').number_format = '$#,##0'
    ws.cell(row=r, column=9, value=round(pulp_ans[name]))
    for j in range(1, 10):
        ws.cell(row=r, column=j).border = box
rl = r0 + len(ca.CHANNELS) - 1
tr = rl + 1
ws.cell(row=tr, column=1, value='TOTAL').font = bold
ws.cell(row=tr, column=2, value=f'=SUM(B{r0}:B{rl})')
ws.cell(row=tr, column=6, value=f'=SUM(F{r0}:F{rl})').font = bold
ws.cell(row=tr, column=8, value=f'=SUMPRODUCT(F{r0}:F{rl},C{r0}:C{rl})').number_format = '$#,##0'
ws.cell(row=tr, column=8).font = Font(bold=True, color='1F6FD1')
ws.cell(row=tr + 2, column=1, value='Objective (maximize):').font = bold; ws.cell(row=tr + 2, column=2, value=f'=H{tr}')

notes = [
    'HOW TO RUN',
    '1. Enable Solver once: Mac Tools > Excel Add-ins > Solver Add-in. Windows File > Options > Add-ins > Go > Solver Add-in.',
    '2. Data > Solver. The dialog is pre-filled:',
    f'   Set Objective: $H${tr}  (To: Max)    By Changing: $F${r0}:$F${rl}',
    f'   Constraints: F <= Demand (B),  F >= Min units (E),  total F (F{tr}) <= Supply (B4)',
    '   Method: Simplex LP,  Make Unconstrained Variables Non-Negative: on',
    '3. Solve > keep solution > also tick "Sensitivity" report.',
    '4. In the Sensitivity report, the Shadow Price on the supply row is the value of ONE more unit of supply;',
    '   shadow prices on the floor rows are what each commitment costs per unit.',
    '5. Compare column F with column I (same model solved in Python with PuLP).',
]
for k, t in enumerate(notes):
    ws.cell(row=tr + 4 + k, column=1, value=t).font = bold if k == 0 else Font()
for col, w in zip('ABCDEFGHI', (22, 10, 14, 10, 11, 13, 9, 12, 12)):
    ws.column_dimensions[col].width = w

# --- pre-load Solver settings (Solver reads these sheet-scoped hidden names) ---
S = "'Allocation'!"
names = {
    'solver_opt': f'{S}$H${tr}', 'solver_typ': '1', 'solver_val': '0',
    'solver_adj': f'{S}$F${r0}:$F${rl}', 'solver_eng': '2', 'solver_neg': '1', 'solver_ver': '3',
    'solver_num': '3',
    'solver_lhs1': f'{S}$F${r0}:$F${rl}', 'solver_rel1': '1', 'solver_rhs1': f'{S}$B${r0}:$B${rl}',
    'solver_lhs2': f'{S}$F${r0}:$F${rl}', 'solver_rel2': '3', 'solver_rhs2': f'{S}$E${r0}:$E${rl}',
    'solver_lhs3': f'{S}$F${tr}', 'solver_rel3': '1', 'solver_rhs3': f'{S}$B$4',
}
for k, v in names.items():
    dn = DefinedName(k, attr_text=v, hidden=True)
    ws.defined_names[k] = dn

out = os.path.join(HERE, 'constrained_allocation_solver.xlsx')
wb.save(out)
print('wrote', out)
